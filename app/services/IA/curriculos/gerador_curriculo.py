from pathlib import Path

from docx import Document
from docx.shared import Pt
from google.genai import types

from app.services.IA.analise.analise_service import AnaliseService


class GeradorCurriculo:
    """Personaliza o currículo usando somente fatos presentes na base."""

    def __init__(self):
        self.analise_service = AnaliseService()

    def gerar_texto(self, vaga, analise):
        prompt = f"""Personalize o currículo-base para a vaga descrita abaixo.
            Use português brasileiro e mantenha somente fatos explicitamente presentes
            no currículo-base. Não invente experiências, tecnologias, resultados, datas,
            certificações ou métricas. Não copie requisitos da vaga como se fossem
            competências do candidato. Omita informações ausentes. Retorne o currículo
            em texto simples, com seções e bullets claros, sem comentários sobre o processo.

            CURRÍCULO-BASE:
            {self.analise_service.curriculo}

            ANÁLISE DA VAGA:
            {analise}

            DADOS DA VAGA:
            Cargo: {vaga.titulo}
            Empresa: {vaga.empresa or 'Não informada'}
            Localidade: {vaga.localidade or 'Não informada'}
            Descrição: {vaga.descricao or 'Não informada'}
            """
        response = self.analise_service.client.models.generate_content(
            model="gemini-2.5-flash",
            contents=[prompt],
            config=types.GenerateContentConfig(temperature=0.2),
        )
        if not response.text or not response.text.strip():
            raise ValueError("A IA não retornou o currículo.")
        return response.text.strip()

    @staticmethod
    def criar_docx(texto, caminho):
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Pt(45)
        section.bottom_margin = Pt(45)
        for linha in texto.splitlines():
            linha = linha.strip()
            if not linha:
                continue
            if linha.startswith("# "):
                doc.add_heading(linha[2:], level=1)
            elif linha.startswith("## "):
                doc.add_heading(linha[3:], level=2)
            elif linha.startswith(("- ", "• ")):
                doc.add_paragraph(linha[2:], style="List Bullet")
            else:
                doc.add_paragraph(linha)
        doc.save(Path(caminho))
